import os
import re
import sys
from docx import Document
from docx.shared import Pt

# Mapping between placeholders and sections
PLACEHOLDER_MAP = {
    "<< EXEC_SUMMARY >>": "Executive Summary",
    "<< KEY_DISCUSSION >>": "Key Discussion Points",
    "<< ACTION_ITEMS >>": "Action Items and Next Steps",
    "<< DETAILED_SUMMARY >>": "Detailed Session Summary",
    "<< APPENDIX >>": "Appendix"
}

def clean_text(text):
    # Remove control characters not compatible with XML and ANSI codes (e.g., [0m)
    text = re.sub(r'\[\d+m', '', text)
    text = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', '', text)
    return text.strip()

def extract_sections(text):
    # Match both # Executive Summary and **Executive Summary**
    pattern = r"(?:#|\*\*)\s*(Executive Summary|Key Discussion Points|Action Items and Next Steps|Detailed Session Summary)\s*(?:\*\*)?"
    matches = list(re.finditer(pattern, text))
    sections = {section: [] for section in PLACEHOLDER_MAP.values() if section != "Appendix"}

    for i, match in enumerate(matches):
        section_name = match.group(1)
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        content = text[start:end].strip()
        sections[section_name].append(clean_text(content))

    return sections

def replace_placeholders(document, content_map):
    for para in document.paragraphs:
        for placeholder, section_name in PLACEHOLDER_MAP.items():
            if placeholder in para.text:
                full_text = "\n\n".join(content_map.get(section_name, [""]))
                replace_text_in_paragraph(para, placeholder, full_text)

def replace_text_in_paragraph(paragraph, placeholder, replacement):
    for run in paragraph.runs:
        if placeholder in run.text:
            run.text = run.text.replace(placeholder, replacement)

def create_report(input_dir, base_filename, output_docx, template_path):
    document = Document(template_path)
    all_sections = {section: [] for section in PLACEHOLDER_MAP.values()}

    part_num = 1
    while True:
        part_path = os.path.join(input_dir, f"{base_filename}_part{part_num}_output.txt")
        if not os.path.exists(part_path):
            break

        with open(part_path, "r", encoding="utf-8") as f:
            text = f.read()
            extracted = extract_sections(text)
            for key, value in extracted.items():
                if value:
                    all_sections[key].append(f"Part {part_num}\n" + clean_text("\n".join(value)))
        part_num += 1

    # Debug: Save all_sections content to file before docx output
    try:
        with open("debug_sections.txt", "w", encoding="utf-8") as debug_f:
            for section, contents in all_sections.items():
                debug_f.write(f"=== {section} ===\n")
                debug_f.write("\n\n".join(contents))
                debug_f.write("\n\n")
    except Exception as e:
        print(f"[Warning] Failed to write debug_sections.txt: {e}")

    replace_placeholders(document, all_sections)
    document.save(output_docx)
    print(f"✅ Final report saved to: {output_docx}")

if __name__ == "__main__":
    if len(sys.argv) != 5:
        print("Usage: python create_docx.py <input_dir> <base_filename> <output_file> <template_file>")
        sys.exit(1)

    input_dir = sys.argv[1]
    base_filename = sys.argv[2]
    output_file = sys.argv[3]
    template_file = sys.argv[4]

    create_report(input_dir, base_filename, output_file, template_file)
